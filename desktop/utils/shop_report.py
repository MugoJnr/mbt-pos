"""Full shop spreadsheet. Only sheets for tables that exist. No secrets."""
from __future__ import annotations

from datetime import datetime, timedelta
from io import BytesIO

_SECRET_PARTS = ('password', 'hash', 'token', 'secret', 'pin', 'api_key', 'jwt')
_ROW_CAP = 50000

_SHEETS = (
    ('Sales', 'sales', (
        'id', 'receipt_number', 'cashier_name', 'subtotal', 'discount', 'tax',
        'total', 'payment_method', 'amount_paid', 'status', 'created_at',
    ), 'created_at'),
    ('Sales Items', 'sale_items', (
        'id', 'sale_id', 'product_name', 'sku', 'quantity', 'unit_price',
        'unit_cost', 'discount', 'total',
    ), None),
    ('Inventory', 'products', (
        'id', 'name', 'sku', 'barcode', 'category', 'unit', 'price', 'cost_price',
        'stock', 'min_stock', 'is_active',
    ), None),
    ('Stock Movements', 'stock_movements', (
        'id', 'product_name', 'movement_type', 'qty_before', 'qty_change',
        'qty_after', 'reference', 'reason', 'username', 'created_at',
    ), 'created_at'),
    ('Purchases', 'purchases', (
        'id', 'purchase_number', 'supplier_name', 'total', 'status',
        'payment_method', 'delivery_date', 'received_by_name', 'created_at',
    ), 'created_at'),
    ('Debts', 'debt_invoices', (
        'id', 'invoice_number', 'customer_name', 'total_amount', 'amount_paid',
        'balance', 'status', 'created_at',
    ), 'created_at'),
    ('Debt Payments', 'debt_payments', (
        'id', 'payment_receipt', 'invoice_id', 'amount', 'payment_method',
        'balance_after', 'cashier_name', 'created_at',
    ), 'created_at'),
    ('Expenses', 'expenses', (
        'id', 'category', 'amount', 'description', 'payment_method',
        'cashier_name', 'status', 'created_at',
    ), 'created_at'),
    ('Internal Consumption', 'stock_consumptions', (
        'id', 'reference_no', 'date', 'reason', 'notes', 'created_at',
    ), 'created_at'),
    ('Stocktakes', 'stocktakes', (
        'id', 'reference', 'name', 'status', 'reason', 'counting_mode',
        'created_by_name', 'started_at', 'submitted_at', 'completed_at',
    ), 'started_at'),
    ('Users', 'users', (
        'id', 'username', 'full_name', 'role', 'is_active', 'last_login',
    ), None),
    ('Audit Log', 'audit_log', (
        'id', 'username', 'action', 'module', 'details', 'created_at',
    ), 'created_at'),
)


def period_bounds(preset: str, start: str = '', end: str = '') -> tuple[str, str, str]:
    today = datetime.now().date()
    key = (preset or 'today').strip().lower().replace(' ', '_')
    if key in ('all', 'all_available_data'):
        return 'All available data', '', ''
    if key == 'yesterday':
        day = today - timedelta(days=1)
        return 'Yesterday', day.isoformat(), day.isoformat()
    if key in ('this_week', 'week'):
        start_day = today - timedelta(days=today.weekday())
        return 'This week', start_day.isoformat(), today.isoformat()
    if key in ('this_month', 'month'):
        return 'This month', today.replace(day=1).isoformat(), today.isoformat()
    if key == 'last_month':
        first_this = today.replace(day=1)
        last = first_this - timedelta(days=1)
        return 'Last month', last.replace(day=1).isoformat(), last.isoformat()
    if key in ('custom', 'custom_date_range'):
        return 'Custom', (start or '')[:10], (end or '')[:10]
    return 'Today', today.isoformat(), today.isoformat()


def _table_columns(conn, table: str) -> set[str]:
    rows = conn.execute(f"PRAGMA table_info({table})").fetchall()
    return {row[1] for row in rows}


def _table_exists(conn, table: str) -> bool:
    row = conn.execute(
        "SELECT 1 FROM sqlite_master WHERE type='table' AND name=?", (table,)
    ).fetchone()
    return bool(row)


def build_shop_workbook(conn, *, shop_name: str, preset: str, start: str, end: str,
                        generated_by: str) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    label, start_day, end_day = period_bounds(preset, start, end)
    wb = Workbook()
    header_font = Font(bold=True, color='FFFFFF')
    fill = PatternFill('solid', fgColor='1F4E79')

    def write(title, headers, rows):
        ws = wb.active if wb.active.title == 'Sheet' else wb.create_sheet()
        ws.title = title[:31]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(1, col, header)
            cell.font = header_font
            cell.fill = fill
        for r, row in enumerate(rows, 2):
            for c, value in enumerate(row, 1):
                ws.cell(r, c, value)
        if headers:
            ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(1, len(rows)+1)}"
            ws.freeze_panes = 'A2'
            for col in range(1, len(headers) + 1):
                ws.column_dimensions[get_column_letter(col)].width = 18
        return ws

    sales_total = sales_count = 0
    if _table_exists(conn, 'sales'):
        where, params = _date_clause('created_at', start_day, end_day)
        row = conn.execute(
            f"SELECT COUNT(*), COALESCE(SUM(CASE WHEN status='voided' THEN 0 ELSE total END),0) "
            f"FROM sales {where}",
            params,
        ).fetchone()
        sales_count = int(row[0] or 0)
        sales_total = round(float(row[1] or 0), 2)
    write('Shop Summary', ['Field', 'Value'], [
        ['Shop', shop_name or 'Shop'],
        ['Report period', label],
        ['From', start_day or 'beginning'],
        ['To', end_day or 'latest'],
        ['Generated at', datetime.now().strftime('%Y-%m-%d %H:%M:%S')],
        ['Generated by', generated_by or ''],
        ['Sales count', sales_count],
        ['Sales total', sales_total],
        ['Note', 'Figures use this shop database. Sheets skip tables that are not in use.'],
    ])
    notes = []
    for title, table, columns, date_col in _SHEETS:
        if not _table_exists(conn, table):
            continue
        present = _table_columns(conn, table)
        safe = [
            col for col in columns
            if col in present and not any(part in col.lower() for part in _SECRET_PARTS)
        ]
        if not safe:
            continue
        where, params = '', []
        if table == 'sale_items' and _table_exists(conn, 'sales') and (start_day or end_day):
            where, params = _date_clause('s.created_at', start_day, end_day)
            sql = (
                f"SELECT {', '.join('i.' + col for col in safe)} FROM sale_items i "
                f"JOIN sales s ON s.id=i.sale_id {where} ORDER BY i.id DESC LIMIT ?"
            )
        else:
            if date_col and date_col in present and (start_day or end_day):
                where, params = _date_clause(date_col, start_day, end_day)
            sql = f"SELECT {', '.join(safe)} FROM {table} {where} ORDER BY id DESC LIMIT ?"
        fetched = conn.execute(sql, params + [_ROW_CAP + 1]).fetchall()
        truncated = len(fetched) > _ROW_CAP
        fetched = fetched[:_ROW_CAP]
        write(title, safe, [list(row) for row in fetched])
        if truncated:
            notes.append(f'{title} stopped at {_ROW_CAP} rows')
    if notes:
        write('Export Notes', ['Note'], [[note] for note in notes])
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _date_clause(column: str, start_day: str, end_day: str):
    clauses, params = [], []
    if start_day:
        clauses.append(f"substr({column},1,10)>=?")
        params.append(start_day)
    if end_day:
        clauses.append(f"substr({column},1,10)<=?")
        params.append(end_day)
    if not clauses:
        return '', []
    return 'WHERE ' + ' AND '.join(clauses), params
