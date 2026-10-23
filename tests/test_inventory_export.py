"""Inventory workbook must still exist after export cleanup."""
import os

from openpyxl import load_workbook

from backend.report_export_service import export_inventory_full


def test_inventory_export_keeps_the_workbook(tmp_path, monkeypatch):
    monkeypatch.setattr(
        'backend.report_export_service.get_export_dir',
        lambda: str(tmp_path),
    )
    path = export_inventory_full(
        [{'id': 1, 'name': 'QA Item', 'sku': 'QA', 'category': 'QA',
          'price': 10, 'cost_price': 4, 'stock': 2, 'min_stock': 1}],
        [{'product_id': 1, 'movement_type': 'sale', 'qty_change': -1,
          'qty_before': 3, 'qty_after': 2, 'created_at': '2026-10-16 12:00:00'}],
        shop_name='QA Shop',
        currency='KES',
        generated_by='test',
    )
    assert os.path.isfile(path)
    assert path.startswith(str(tmp_path))
    wb = load_workbook(path)
    assert 'Inventory' in wb.sheetnames
    assert 'Stock Movements' in wb.sheetnames
    leftovers = [name for name in os.listdir(tmp_path) if name.startswith('_tmp_')]
    assert leftovers == []
