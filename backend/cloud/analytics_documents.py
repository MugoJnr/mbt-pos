"""Formatted portal report downloads (Excel, PDF, printable HTML).

Rows come from the cloud analytics tables. This module only lays them out.
It does not read the shop SQLite database.
"""
from __future__ import annotations

import html
import io
from typing import Any

_MONEY = {
    'subtotal', 'discount', 'tax', 'total', 'amount_paid', 'change_amount',
    'total_amount', 'balance', 'amount', 'price', 'cost_price',
}
_QTY = {'stock', 'min_stock', 'quantity', 'qty'}
_LABELS = {
    'receipt_number': 'Receipt',
    'invoice_number': 'Invoice',
    'payment_receipt': 'Payment receipt',
    'created_at': 'Date',
    'source_created_at': 'Date',
    'due_date': 'Due',
    'cashier_name': 'Cashier',
    'customer_name': 'Customer',
    'customer_phone': 'Phone',
    'payment_method': 'Method',
    'amount_paid': 'Paid',
    'change_amount': 'Change',
    'total_amount': 'Original',
    'cost_price': 'Cost',
    'min_stock': 'Minimum',
    'is_active': 'Active',
    'device_id': 'Device',
    'source_id': 'Record',
}


def _label(field: str) -> str:
    return _LABELS.get(field, field.replace('_', ' ').title())


def _kind(field: str) -> str:
    if field in _MONEY:
        return 'currency'
    if field in _QTY:
        return 'qty'
    if field.endswith('_at'):
        return 'datetime'
    if field.endswith('_date') or field == 'due_date':
        return 'date'
    return 'text'


def _money(value: Any) -> str:
    try:
        amount = float(value or 0)
    except (TypeError, ValueError):
        amount = 0.0
    return f'KSh {amount:,.2f}'


def sheet_from_rows(title: str, rows: list[dict], fields: list[str]) -> dict:
    return {
        'title': title[:31] or 'Report',
        'headers': [_label(field) for field in fields],
        'rows': [[row.get(field, '') for field in fields] for row in rows],
        'kinds': [_kind(field) for field in fields],
        'count': len(rows),
    }


def summary_sheet(overview: dict, *, can_see_finance: bool) -> dict:
    summary = overview.get('summary') or {}
    lines = [
        ('Period', f"{overview.get('start') or ''} to {overview.get('end') or ''}"),
        ('Gross sales', _money(summary.get('gross_sales'))),
        ('Collected', _money(summary.get('collected_revenue'))),
        ('Transactions', summary.get('transactions') or 0),
        ('Average sale', _money(summary.get('avg_sale'))),
        ('Discounts', _money(summary.get('discounts'))),
        ('Tax', _money(summary.get('tax'))),
        ('Items sold', summary.get('items_sold') or 0),
        ('Voided sales', summary.get('void_transactions') or 0),
        ('Debt issued', _money(summary.get('debt_issued'))),
        ('Debt collected', _money(summary.get('debt_collected'))),
        ('Debt outstanding', _money(summary.get('debt_outstanding'))),
        ('Overdue invoices', summary.get('overdue_count') or 0),
    ]
    if can_see_finance:
        lines.extend([
            ('Cost of goods', _money(summary.get('cost_of_goods'))),
            ('Gross profit', _money(summary.get('gross_profit'))),
            ('Inventory value', _money(summary.get('inventory_value'))),
        ])
    lines.append(('Last sync', summary.get('last_sync_at') or 'Not recorded'))
    return {
        'title': 'Summary',
        'headers': ['Measure', 'Value'],
        'rows': [[label, value] for label, value in lines],
        'kinds': ['text', 'text'],
        'count': len(lines),
    }


def _drop_cost(fields: list[str], *, can_see_finance: bool) -> list[str]:
    if can_see_finance:
        return fields
    hidden = {'cost_price', 'unit_cost', 'cost', 'gross_profit', 'profit'}
    return [field for field in fields if field not in hidden]


def collect_portal_sheets(
    org_id: str,
    report: str,
    *,
    start: str,
    end: str,
    can_see_finance: bool,
    status: str = '',
    payment: str = '',
    cashier: str = '',
    customer: str = '',
    q: str = '',
    stock: str = '',
    sort: str = '',
    order: str = 'desc',
) -> list[dict]:
    """Build the sheets for one portal download."""
    from backend.cloud.platform_service import analytics_export_rows, analytics_overview

    kind = (report or 'sales').strip().lower()
    filters = dict(
        start=start, end=end, status=status, payment=payment, cashier=cashier,
        customer=customer, q=q, stock=stock, sort=sort, order=order,
    )

    def table(name: str, title: str) -> dict:
        rows, fields = analytics_export_rows(org_id, report=name, **filters)
        return sheet_from_rows(title, rows, _drop_cost(fields, can_see_finance=can_see_finance))

    if kind == 'shop':
        overview = analytics_overview(org_id, start=start, end=end)
        return [
            summary_sheet(overview, can_see_finance=can_see_finance),
            table('sales', 'Sales'),
            table('debts', 'Debt invoices'),
            table('debt_payments', 'Payments'),
            table('inventory', 'Inventory'),
        ]
    if kind == 'overview':
        overview = analytics_overview(org_id, start=start, end=end)
        products = [
            {
                'name': item.get('name'),
                'category': item.get('category'),
                'qty': item.get('qty'),
                'revenue': item.get('revenue'),
            }
            for item in (overview.get('top_products') or [])
        ]
        methods = overview.get('payment_methods') or overview.get('payment_mix') or []
        return [
            summary_sheet(overview, can_see_finance=can_see_finance),
            sheet_from_rows('Top products', products, ['name', 'category', 'qty', 'revenue']),
            sheet_from_rows(
                'Payments',
                list(methods),
                [key for key in ('payment_method', 'method', 'count', 'total', 'amount')
                 if any(key in row for row in methods)] or ['payment_method', 'total'],
            ),
        ]
    if kind in ('debts', 'debt', 'debt_invoices'):
        return [table('debts', 'Debt invoices'), table('debt_payments', 'Payments')]
    if kind in ('debt_payments', 'payments'):
        return [table('debt_payments', 'Payments')]
    if kind in ('inventory', 'products', 'stock'):
        return [table('inventory', 'Inventory')]
    return [table('sales', 'Sales')]


def render_portal_xlsx(
    sheets: list[dict],
    *,
    shop_name: str,
    title: str,
    period: str,
    generated_by: str = '',
    filters: str = '',
) -> bytes:
    from openpyxl.utils import get_column_letter
    from backend.report_export_service import (
        apply_data_cell,
        finalize_table,
        new_workbook_sheet,
        style_header_row,
        write_footer,
        write_report_header,
    )

    first = sheets[0] if sheets else {
        'title': 'Report', 'headers': ['Note'], 'rows': [['No rows']], 'kinds': ['text'],
    }
    wb, first_ws = new_workbook_sheet(first['title'])
    pages = [(first_ws, first)]
    for extra in sheets[1:]:
        pages.append((wb.create_sheet(extra['title'][:31] or 'Sheet'), extra))
    for ws, sheet in pages:
        headers = sheet['headers'] or ['Note']
        rows = sheet['rows']
        kinds = list(sheet.get('kinds') or ['text'] * len(headers))
        header_row = write_report_header(
            ws,
            shop_name=shop_name or 'MBT POS',
            title=title,
            ncols=len(headers),
            period=period,
            generated_by=generated_by or 'Portal',
            filters=filters or sheet['title'],
            currency='KES',
        )
        style_header_row(ws, header_row, headers)
        for index, row in enumerate(rows):
            excel_row = header_row + 1 + index
            for column, kind in enumerate(kinds):
                value = row[column] if column < len(row) else None
                apply_data_cell(
                    ws.cell(row=excel_row, column=column + 1),
                    value,
                    kind=kind,
                    currency='KES',
                    alt=index % 2 == 1,
                )
        data_end = header_row + len(rows)
        write_footer(
            ws,
            (data_end if rows else header_row) + 2,
            len(headers),
            record_count=len(rows),
        )
        if rows:
            finalize_table(ws, header_row, header_row + 1, data_end, len(headers))
        else:
            ws.freeze_panes = f'A{header_row + 1}'
        for column, header in enumerate(headers, start=1):
            letter = get_column_letter(column)
            current = ws.column_dimensions[letter].width or 0
            ws.column_dimensions[letter].width = max(current, min(28, len(str(header)) + 4))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def render_portal_html(
    sheets: list[dict],
    *,
    shop_name: str,
    title: str,
    period: str,
    generated_by: str = '',
    auto_print: bool = False,
) -> str:
    sections = []
    for sheet in sheets:
        head = ''.join(f'<th>{html.escape(str(name))}</th>' for name in sheet['headers'])
        body_rows = []
        for row in sheet['rows'][:2000]:
            cells = ''.join(
                f'<td>{html.escape("" if value is None else str(value))}</td>'
                for value in row
            )
            body_rows.append(f'<tr>{cells}</tr>')
        if not body_rows:
            span = max(1, len(sheet['headers']))
            body_rows.append(f'<tr><td colspan="{span}">No rows in this period</td></tr>')
        sections.append(
            f'<h2>{html.escape(sheet["title"])}</h2>'
            f'<table><thead><tr>{head}</tr></thead><tbody>{"".join(body_rows)}</tbody></table>'
        )
    printed = 'window.print();' if auto_print else ''
    return f"""<!DOCTYPE html><html><head><meta charset="utf-8">
<title>{html.escape(title)} — {html.escape(shop_name or 'MBT POS')}</title>
<style>
body{{font-family:Segoe UI,sans-serif;background:#0f172a;color:#e2e8f0;padding:24px}}
h1{{margin:0;color:#f8fafc}} h2{{margin:28px 0 8px;font-size:16px}}
table{{width:100%;border-collapse:collapse}}
th,td{{padding:8px 10px;border-bottom:1px solid #334155;font-size:13px;text-align:left}}
th{{color:#94a3b8;font-size:11px;text-transform:uppercase}}
@media print{{body{{background:#fff;color:#111}} th,td{{border-color:#ccc}} th{{color:#333}}}}
</style></head><body>
<h1>{html.escape(shop_name or 'MBT POS')}</h1>
<p>{html.escape(title)} · {html.escape(period)} · {html.escape(generated_by or 'Portal')}</p>
{''.join(sections)}
<p style="margin-top:24px;font-size:12px;opacity:.7">MugoByte portal · synced shop data</p>
<script>window.onload=function(){{{printed}}}</script>
</body></html>"""


def render_portal_csv(sheets: list[dict]) -> str:
    import csv
    buf = io.StringIO()
    writer = csv.writer(buf)
    for index, sheet in enumerate(sheets):
        if index:
            writer.writerow([])
        writer.writerow([sheet['title']])
        writer.writerow(sheet['headers'])
        writer.writerows(sheet['rows'])
    return '\ufeff' + buf.getvalue()


def render_portal_pdf(
    sheets: list[dict],
    *,
    shop_name: str,
    title: str,
    period: str,
    generated_by: str = '',
) -> bytes:
    from backend.report_export_service import build_report_pdf

    lines = [f'Prepared by {generated_by or "Portal"}', '']
    for sheet in sheets:
        lines.append(sheet['title'])
        headers = sheet['headers']
        preview = sheet['rows'][:40]
        if sheet['title'] == 'Summary':
            for row in sheet['rows']:
                lines.append(f"  {row[0]}: {row[1] if len(row) > 1 else ''}")
        else:
            lines.append(f"  {sheet.get('count', len(sheet['rows']))} rows")
            for row in preview:
                bits = [
                    str(value)
                    for value in row[:min(4, len(headers))]
                    if value not in (None, '')
                ]
                if bits:
                    lines.append('  ' + ' · '.join(bits)[:110])
            if len(sheet['rows']) > len(preview):
                lines.append('  Further rows are in the Excel download.')
        lines.append('')
    return build_report_pdf(title, shop_name or 'MBT POS', period, lines, 'KES')
